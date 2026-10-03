"""Genera los .docx de envío a la IJC desde los .md del paper (sin pandoc).

Formato: Times New Roman 12, interlineado doble, números de línea continuos y número de página al pie, tablas
editables, figuras al final (una por página, con la leyenda debajo), como pide la guía de "free format" de RMetS.
Uso: python analisis/a_docx.py manuscrito|suplemento   → envio/ijc/manuscript.docx | supporting_information.docx
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SALIDA = {"manuscrito": "envio/ijc/manuscript.docx", "suplemento": "envio/ijc/supporting_information.docx"}


def inline(par, texto):
    """Agrega texto con **negrita**, *cursiva* y [enlaces](url) (el enlace queda como texto)."""
    texto = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", texto)
    for trozo in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
        if not trozo:
            continue
        if trozo.startswith("**"):
            par.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("*"):
            par.add_run(trozo[1:-1]).italic = True
        else:
            par.add_run(trozo)


def numeros_de_linea(doc):
    sect = doc.sections[0]._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    sect.append(ln)


def numero_de_pagina(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = 1
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run()
        if tipo:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tipo)
            r._r.append(f)
        else:
            t = OxmlElement("w:instrText")
            t.text = texto
            r._r.append(t)


def tabla(doc, lineas):
    filas = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lineas if not re.match(r"^\|\s*-", l)]
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = "Table Grid"
    for i, fila in enumerate(filas):
        for j, celda in enumerate(fila):
            par = t.cell(i, j).paragraphs[0]
            par.paragraph_format.line_spacing = 1.0
            inline(par, celda)
            for r in par.runs:
                r.font.size = Pt(9)
                r.bold = r.bold or i == 0
    doc.add_paragraph()


def main(nombre):
    md = open(f"paper/{nombre}.md").read()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.line_spacing = 2.0
    st.paragraph_format.space_after = Pt(0)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2.5)
    for h in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        doc.styles[h].font.name = "Times New Roman"
        doc.styles[h].font.color.rgb = None
    numeros_de_linea(doc)
    numero_de_pagina(doc)
    bloques = re.split(r"\n\s*\n", md.strip())
    leyenda_pendiente = None
    for b in bloques:
        b = b.strip()
        if b.startswith("*Manuscript draft"):
            continue
        if b.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(b[2:].strip())
            r.bold, r.font.size = True, Pt(14)
            continue
        m = re.match(r"^(#{2,3}) (.*)", b)
        if m:
            nivel = len(m.group(1)) - 1
            if m.group(2).strip() in ("Figures", "Supporting figures", "Abstract"):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            h = doc.add_heading(m.group(2).strip(), level=nivel)
            for r in h.runs:
                r.font.name, r.font.size = "Times New Roman", Pt(13 if nivel == 1 else 12)
            continue
        if b.startswith("|"):
            tabla(doc, b.splitlines())
            continue
        img = re.match(r"^!\[[^\]]*\]\(([^)]+)\)$", b)
        if img:
            ruta = os.path.normpath(os.path.join("paper", img.group(1)))
            doc.add_picture(ruta, width=Cm(16))
            if leyenda_pendiente:
                p = doc.add_paragraph()
                inline(p, leyenda_pendiente)
                leyenda_pendiente = None
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue
        if re.match(r"^\*\*(Figure|Figure S)\s?S?\d", b):
            leyenda_pendiente = b               # la leyenda va debajo de la figura
            continue
        p = doc.add_paragraph()
        if not b.startswith("**"):                 # sangría en párrafos de texto, no en rótulos
            p.paragraph_format.first_line_indent = Cm(0.75)
        inline(p, " ".join(l.strip() for l in b.splitlines()))
    os.makedirs("envio/ijc", exist_ok=True)
    doc.save(SALIDA[nombre])
    print("ok", SALIDA[nombre])


if __name__ == "__main__":
    os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    main(sys.argv[1] if len(sys.argv) > 1 else "manuscrito")
